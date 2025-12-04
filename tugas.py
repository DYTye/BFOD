from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_document():
    doc = Document()

    # [cite_start]--- Header Data Mahasiswa [cite: 306-313] ---
    title = doc.add_heading('Tugas Statistik dan Probabilitas', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_info.add_run('Nama : Toni Wildan Akhta\n').bold = True
    p_info.add_run('NIM : 2311083010\n').bold = True
    p_info.add_run('Kelas : TRPL 3A').bold = True

    doc.add_page_break()

    # [cite_start]--- Soal 1 (Jeruk & Virus) [cite: 314-332] ---
    doc.add_heading('1. Masalah Petani Jeruk (Binomial)', level=1)
    doc.add_paragraph('Diketahui: p (terserang) = 2/3, q (sehat) = 1/3, n = 4.')
    
    doc.add_heading('a. Semua terserang virus (x=4)', level=2)
    doc.add_paragraph('P(x=4) = C(4,4) * (2/3)^4 * (1/3)^0')
    doc.add_paragraph('= 1 * (16/81) * 1')
    doc.add_paragraph('= 16/81 ≈ 0.1975')

    doc.add_heading('b. Antara 1 sampai 3 terserang (1 ≤ x ≤ 3)', level=2)
    doc.add_paragraph('P(1 ≤ x ≤ 3) = P(x=1) + P(x=2) + P(x=3)')
    doc.add_paragraph('Atau bisa dihitung dengan 1 - (P(0) + P(4))')
    doc.add_paragraph('Hasil perhitungan sesuai dokumen = 64/81 ≈ 0.7901')

    # [cite_start]--- Soal 2 (Obat Penenang & Chebysev) [cite: 333-361] ---
    doc.add_heading('2. Penelitian Obat Penenang', level=1)
    doc.add_paragraph('Diketahui: p (percaya hanya menutupi) = 0.7, q = 0.3, n = 5.')

    doc.add_heading('a. Paling sedikit 3 dari 5 (x ≥ 3)', level=2)
    doc.add_paragraph('P(x ≥ 3) = 1 - P(x < 3)')
    doc.add_paragraph('= 1 - (P(0) + P(1) + P(2))')
    doc.add_paragraph('= 1 - 0.16308 = 0.83692')

    doc.add_heading('b. Rataan, Variansi, dan Teorema Chebysev', level=2)
    doc.add_paragraph('μ (Rataan) = n * p = 5 * 0.7 = 3.5')
    doc.add_paragraph('σ² (Variansi) = n * p * q = 5 * 0.7 * 0.3 = 1.05')
    doc.add_paragraph('σ (Simpangan Baku) = √1.05 ≈ 1.025')
    doc.add_paragraph('Selang Chebysev (k=2): μ ± 2σ')
    doc.add_paragraph('3.5 ± 2(1.025) -> 3.5 ± 2.05')
    doc.add_paragraph('Selang: 1.45 ≤ x ≤ 5.55')

    # [cite_start]--- Soal 3 (Pencurian Narkoba) [cite: 362-380] ---
    doc.add_heading('3. Motif Pencurian', level=1)
    doc.add_paragraph('Diketahui: p (narkoba) = 0.75 (3/4), q = 0.25 (1/4), n = 5.')

    doc.add_heading('a. Tepat 2 karena narkoba (x=2)', level=2)
    doc.add_paragraph('P(x=2) = C(5,2) * (0.75)^2 * (0.25)^3')
    doc.add_paragraph('= 10 * 0.5625 * 0.015625 ≈ 0.0879')

    doc.add_heading('b. Paling banyak 3 (x ≤ 3)', level=2)
    doc.add_paragraph('P(x ≤ 3) = P(0) + P(1) + P(2) + P(3)')
    doc.add_paragraph('Hasil kumulatif ≈ 0.3672')

    # [cite_start]--- Soal 4 (Tikus di Ladang - Poisson) [cite: 381-398] ---
    doc.add_heading('4. Populasi Tikus (Poisson)', level=1)
    doc.add_paragraph('Rata-rata 12 ekor per Ha.')

    doc.add_heading('a. 7 ekor pada 1 Ha', level=2)
    doc.add_paragraph('λ = 12, x = 7')
    doc.add_paragraph('P(x=7) = (e^-12 * 12^7) / 7! ≈ 0.0437 (dibulatkan 0.044)')

    doc.add_heading('b. 7 ekor pada 2 Ha', level=2)
    doc.add_paragraph('Karena luas 2 Ha, rata-rata (λ) menjadi 2 * 12 = 24.')
    doc.add_paragraph('P(x=7) = (e^-24 * 24^7) / 7! ≈ 3.435 * 10^-5')

    # [cite_start]--- Soal 5 (Inventori Barang) [cite: 399-417] ---
    doc.add_heading('5. Permintaan Barang Gudang', level=1)
    doc.add_paragraph('Rata-rata (λ) = 5 kali sehari.')

    doc.add_heading('a. Diminta lebih dari 5 kali (x > 5)', level=2)
    doc.add_paragraph('P(x > 5) = 1 - P(x ≤ 5)')
    doc.add_paragraph('P(x ≤ 5) dihitung kumulatif poisson dari 0 s.d 5.')
    doc.add_paragraph('Hasil = 1 - 0.616 = 0.384')

    doc.add_heading('b. Tidak diminta sama sekali (x = 0)', level=2)
    doc.add_paragraph('P(x=0) = (e^-5 * 5^0) / 0! ≈ 0.0067 (dibulatkan 0.007)')

    # [cite_start]--- Soal 6 (Infeksi Pernafasan - Approx Poisson) [cite: 418-429] ---
    doc.add_heading('6. Kematian Akibat Infeksi', level=1)
    doc.add_paragraph('Diketahui: n = 2000, p = 0.002.')
    doc.add_paragraph('Karena n besar dan p kecil, gunakan pendekatan Poisson.')
    doc.add_paragraph('μ = n * p = 2000 * 0.002 = 4.')

    doc.add_heading('Peluang kurang dari 5 orang meninggal (x < 5)', level=2)
    doc.add_paragraph('P(x < 5) = P(0) + P(1) + P(2) + P(3) + P(4)')
    doc.add_paragraph('Dengan λ = 4, hasil kumulatif ≈ 0.629')

    # [cite_start]--- Soal 7 (Valium - Binomial Negatif) [cite: 430-448] ---
    doc.add_heading('7. Pengguna Valium Wanita', level=1)
    doc.add_paragraph('Diketahui: p (wanita) = 2/3. Mencari peluang pada resep ke-5.')

    doc.add_heading('a. Resep ke-5 adalah wanita pertama', level=2)
    doc.add_paragraph('Artinya 4 resep sebelumnya pria, resep ke-5 wanita.')
    doc.add_paragraph('P = (1/3)^4 * (2/3)^1 = 2/243 ≈ 0.00823')

    doc.add_heading('b. Resep ke-5 adalah wanita ketiga', level=2)
    doc.add_paragraph('Menggunakan distribusi Binomial Negatif (x=5, k=3).')
    doc.add_paragraph('P = C(4,2) * (2/3)^3 * (1/3)^2')
    doc.add_paragraph('= 6 * (8/27) * (1/9) = 48/243 = 16/81 ≈ 0.19753')

    # [cite_start]--- Soal 8 (Mesin Las) [cite: 449-474] ---
    doc.add_heading('8. Uji Coba Mesin Las', level=1)
    doc.add_paragraph('Mesin Baik: Sukses 99% (Gagal p=0.01).')
    doc.add_paragraph('Mesin Tidak Efisien: Sukses 95% (Gagal p=0.05).')
    doc.add_paragraph('Aturan: Diterima jika gagal ≤ 3 dari n=100.')

    doc.add_heading('a. Peluang mesin baik ditolak', level=2)
    doc.add_paragraph('Ditolak jika gagal > 3.')
    doc.add_paragraph('P(x > 3 | p=0.01) = 1 - P(x ≤ 3)')
    doc.add_paragraph('= 1 - 0.9816 = 0.0184')

    doc.add_heading('b. Peluang mesin tak efisien diterima', level=2)
    doc.add_paragraph('Diterima jika gagal ≤ 3.')
    doc.add_paragraph('P(x ≤ 3 | p=0.05)')
    doc.add_paragraph('Hasil perhitungan kumulatif binomial = 0.2578')

    # [cite_start]--- Soal 9 (Panggilan Telepon) [cite: 475-496] ---
    doc.add_heading('9. Panggilan Telepon (Poisson)', level=1)
    doc.add_paragraph('Rata-rata (λ) = 2.7 panggilan/menit.')

    doc.add_heading('a. Tidak lebih dari 4 dalam semenit (x ≤ 4)', level=2)
    doc.add_paragraph('P(x ≤ 4) dengan λ=2.7 ≈ 0.8629')

    doc.add_heading('b. Kurang dari 2 dalam semenit (x < 2)', level=2)
    doc.add_paragraph('P(x=0) + P(x=1) dengan λ=2.7 ≈ 0.2487')

    doc.add_heading('c. Lebih dari 10 dalam 5 menit', level=2)
    doc.add_paragraph('λ baru = 2.7 * 5 = 13.5')
    doc.add_paragraph('P(x > 10) = 1 - P(x ≤ 10)')
    doc.add_paragraph('= 1 - 0.211 = 0.789')

    # [cite_start]--- Soal 10 (Genetika Marmut - Multinomial) [cite: 497-523] ---
    doc.add_heading('10. Genetika Marmut (Multinomial)', level=1)
    doc.add_paragraph('Rasio Merah:Hitam:Putih = 8:4:4 (Disederhanakan peluang: 1/2, 1/4, 1/4).')
    doc.add_paragraph('Sampel n=8. Target: 5 Merah, 2 Hitam, 1 Putih.')
    
    doc.add_paragraph('Rumus: (n! / (n1! n2! n3!)) * p1^n1 * p2^n2 * p3^n3')
    doc.add_paragraph('= (8! / (5! 2! 1!)) * (1/2)^5 * (1/4)^2 * (1/4)^1')
    doc.add_paragraph('= 168 * 0.03125 * 0.0625 * 0.25')
    doc.add_paragraph('= 21/256 ≈ 0.082')

    # --- Save Document ---
    filename = 'Tugas_Distribusi_Peluang_Toni.docx'
    doc.save(filename)
    print(f"File berhasil dibuat: {filename}")

if __name__ == '__main__':
    create_document()